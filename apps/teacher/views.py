import io
import datetime

from django.shortcuts import redirect, render, get_object_or_404
from django.contrib.auth import logout
from django.http import HttpResponse
from django.contrib import messages

from apps.accounts.decorators import role_required
from apps.ai_engine.models import LessonPlan
from apps.ai_engine.services import ai_service


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------

@role_required("teacher")
def dashboard(request):
    return render(request, "teacher/dashboard.html")


# ---------------------------------------------------------------------------
# Lesson Planner — Django-native form handling (no JS API calls)
# ---------------------------------------------------------------------------

@role_required("teacher")
def lesson_planner(request):
    generated_plan = None

    if request.method == "POST":
        title      = request.POST.get("title",      "").strip()
        subject    = request.POST.get("subject",    "").strip()
        grade      = request.POST.get("grade",      "").strip()
        duration   = request.POST.get("duration",   "45").strip()
        objectives = request.POST.get("objectives", "").strip()
        curriculum = request.POST.get("curriculum", "").strip() or "Standard Curriculum"

        if not all([title, subject, grade, duration, objectives]):
            messages.error(request, "Please fill in all required fields.")
        else:
            try:
                duration_int = int(duration)
                content = ai_service.generate_lesson_plan(
                    subject=subject,
                    grade=grade,
                    duration=duration_int,
                    objectives=objectives,
                    curriculum=curriculum,
                )
                plan = LessonPlan.objects.create(
                    teacher=request.user,
                    school=getattr(request.user, "school", None),
                    title=title,
                    subject=subject,
                    grade=grade,
                    duration=duration_int,
                    objectives=objectives,
                    content=content,
                )
                generated_plan = plan
                messages.success(request, "Lesson plan generated successfully!")
            except Exception as e:
                messages.error(request, f"Failed to generate lesson plan: {e}")

    previous_plans = LessonPlan.objects.filter(
        teacher=request.user
    ).order_by("-created_at")[:10]

    return render(request, "teacher/lesson_planner.html", {
        "generated_plan": generated_plan,
        "previous_plans": previous_plans,
    })


# ---------------------------------------------------------------------------
# Download — streams Word or PDF directly
# ---------------------------------------------------------------------------

@role_required("teacher")
def lesson_plan_download(request, pk, fmt):
    plan = get_object_or_404(LessonPlan, pk=pk, teacher=request.user)
    if fmt == "docx":
        return _build_docx(plan)
    if fmt == "pdf":
        return _build_pdf(plan)
    messages.error(request, "Invalid download format.")
    return redirect("lesson_planner")


# ===========================================================================
#  Word document builder — professional, submission-ready
# ===========================================================================

def _build_docx(plan):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    teacher      = plan.teacher
    school       = getattr(teacher, "school", None)
    school_name  = school.name if school else "School"
    teacher_name = teacher.get_full_name() or teacher.username
    c = plan.content if isinstance(plan.content, dict) else {}

    INDIGO     = RGBColor(0x4F, 0x46, 0xE5)
    DARK_BLUE  = RGBColor(0x1E, 0x1B, 0x4B)
    MID_GRAY   = RGBColor(0x6B, 0x72, 0x80)
    WHITE      = RGBColor(0xFF, 0xFF, 0xFF)

    def cell_bg(cell, hex_color):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    def para_underline(p, color_hex="4F46E5", sz=6):
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        btm = OxmlElement("w:bottom")
        btm.set(qn("w:val"), "single")
        btm.set(qn("w:sz"), str(sz))
        btm.set(qn("w:color"), color_hex)
        pBdr.append(btm)
        pPr.append(pBdr)

    def cell_top_border(cell, color_hex="1E1B4B", sz=8):
        tcPr = cell._tc.get_or_add_tcPr()
        tcBdr = OxmlElement("w:tcBorders")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), str(sz))
        top.set(qn("w:color"), color_hex)
        tcBdr.append(top)
        tcPr.append(tcBdr)

    def section_heading(doc, text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text.upper())
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.bold = True
        run.font.color.rgb = INDIGO
        para_underline(p)
        return p

    def body_p(doc, text, size=10.5, bold=False, color=None, indent_cm=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        if indent_cm:
            p.paragraph_format.left_indent = Cm(indent_cm)
        run = p.add_run(str(text))
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = color
        return p

    def bullet_p(doc, text, size=10.5):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0.63)
        run = p.add_run(str(text))
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        return p

    # ── Document setup ────────────────────────────────────────────────────────
    doc = Document()
    sec = doc.sections[0]
    sec.page_width    = Cm(21)
    sec.page_height   = Cm(29.7)
    sec.left_margin   = Cm(2.54)
    sec.right_margin  = Cm(2.54)
    sec.top_margin    = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    # ── 1. Header banner ─────────────────────────────────────────────────────
    banner = doc.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    bc = banner.rows[0].cells[0]
    cell_bg(bc, "1E1B4B")
    p1 = bc.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(10)
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run(school_name.upper())
    r1.font.name = "Calibri"; r1.font.size = Pt(14); r1.bold = True
    r1.font.color.rgb = WHITE
    p2 = bc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(10)
    r2 = p2.add_run("LESSON PLAN")
    r2.font.name = "Calibri"; r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(0xC7, 0xD2, 0xFE)
    doc.add_paragraph()

    # ── 2. Lesson title ───────────────────────────────────────────────────────
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.space_before = Pt(4)
    tp.paragraph_format.space_after = Pt(8)
    tr_ = tp.add_run(plan.title)
    tr_.font.name = "Calibri"; tr_.font.size = Pt(18); tr_.bold = True
    tr_.font.color.rgb = DARK_BLUE

    # ── 3. Metadata table row 1 ───────────────────────────────────────────────
    meta1 = doc.add_table(rows=2, cols=4)
    meta1.style = "Table Grid"
    meta1.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (h, v) in enumerate(zip(
        ["Subject", "Grade / Class", "Duration", "Date Prepared"],
        [plan.subject, plan.grade,
         f"{plan.duration} minutes",
         plan.created_at.strftime("%d %B %Y")],
    )):
        hc = meta1.rows[0].cells[i]; vc = meta1.rows[1].cells[i]
        cell_bg(hc, "4F46E5"); cell_bg(vc, "EEF2FF")
        hp = hc.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hr_ = hp.add_run(h)
        hr_.font.name = "Calibri"; hr_.font.size = Pt(8.5); hr_.bold = True
        hr_.font.color.rgb = WHITE
        vp = vc.paragraphs[0]; vp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        vr_ = vp.add_run(v)
        vr_.font.name = "Calibri"; vr_.font.size = Pt(10); vr_.bold = True
        vr_.font.color.rgb = DARK_BLUE

    # Metadata table row 2
    meta2 = doc.add_table(rows=2, cols=3)
    meta2.style = "Table Grid"
    meta2.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (h, v) in enumerate(zip(
        ["Teacher", "School / Institution", "Curriculum Framework"],
        [teacher_name, school_name, c.get("curriculum", "Standard Curriculum")],
    )):
        hc = meta2.rows[0].cells[i]; vc = meta2.rows[1].cells[i]
        cell_bg(hc, "1E1B4B"); cell_bg(vc, "EEF2FF")
        hp = hc.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hr_ = hp.add_run(h)
        hr_.font.name = "Calibri"; hr_.font.size = Pt(8.5); hr_.bold = True
        hr_.font.color.rgb = WHITE
        vp = vc.paragraphs[0]; vp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        vr_ = vp.add_run(str(v))
        vr_.font.name = "Calibri"; vr_.font.size = Pt(10)
        vr_.font.color.rgb = DARK_BLUE

    doc.add_paragraph()

    # ── 4. Overview ───────────────────────────────────────────────────────────
    if c.get("overview"):
        section_heading(doc, "1.  Lesson Overview")
        body_p(doc, c["overview"])

    # ── 5. Prior Knowledge ────────────────────────────────────────────────────
    if c.get("prior_knowledge"):
        section_heading(doc, "2.  Prior Knowledge Required")
        body_p(doc, c["prior_knowledge"])

    # ── 6. Learning Objectives ────────────────────────────────────────────────
    section_heading(doc, "3.  Learning Objectives")
    for obj in c.get("objectives", [plan.objectives]):
        bullet_p(doc, obj)

    # ── 7. Success Criteria ───────────────────────────────────────────────────
    if c.get("success_criteria"):
        section_heading(doc, "4.  Success Criteria")
        for sc in c["success_criteria"]:
            bullet_p(doc, sc)

    # ── 8. Key Vocabulary ─────────────────────────────────────────────────────
    if c.get("key_vocabulary"):
        section_heading(doc, "5.  Key Vocabulary")
        body_p(doc, "  •  ".join(c["key_vocabulary"]))

    # ── 9. Materials ──────────────────────────────────────────────────────────
    if c.get("materials"):
        section_heading(doc, "6.  Materials & Resources")
        for mat in c["materials"]:
            bullet_p(doc, mat)

    # ── 10. Lesson Procedure (table) ──────────────────────────────────────────
    sections = c.get("sections", [])
    if sections:
        section_heading(doc, "7.  Lesson Procedure")
        proc = doc.add_table(rows=1, cols=5)
        proc.style = "Table Grid"
        proc.alignment = WD_TABLE_ALIGNMENT.CENTER
        col_widths = [Cm(2.8), Cm(1.6), Cm(4.8), Cm(4.8), Cm(3.4)]
        for i, w in enumerate(col_widths):
            proc.rows[0].cells[i].width = w
        for i, ch in enumerate(
            ["Phase", "Time\n(min)", "Teacher Activities",
             "Student Activities", "Notes"]
        ):
            cell = proc.rows[0].cells[i]
            cell_bg(cell, "4F46E5")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            hp = cell.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            hr_ = hp.add_run(ch)
            hr_.font.name = "Calibri"; hr_.font.size = Pt(8.5); hr_.bold = True
            hr_.font.color.rgb = WHITE

        for idx, s in enumerate(sections):
            bg = "FFFFFF" if idx % 2 == 0 else "EEF2FF"
            row = proc.add_row()
            t_acts = s.get("teacher_activities", s.get("activities", []))
            s_acts = s.get("student_activities", [])
            data = [
                s.get("name", ""),
                str(s.get("duration", "")),
                "\n".join(f"• {a}" for a in t_acts),
                "\n".join(f"• {a}" for a in s_acts),
                s.get("teacher_notes", ""),
            ]
            for i, (cell, text) in enumerate(zip(row.cells, data)):
                cell_bg(cell, bg)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                p = cell.paragraphs[0]
                p.alignment = (WD_ALIGN_PARAGRAPH.CENTER
                               if i == 1 else WD_ALIGN_PARAGRAPH.LEFT)
                run = p.add_run(text)
                run.font.name = "Calibri"; run.font.size = Pt(9)
                if i == 0:
                    run.bold = True; run.font.color.rgb = INDIGO
        doc.add_paragraph()

    # ── 11. Assessment ────────────────────────────────────────────────────────
    if c.get("assessment"):
        section_heading(doc, "8.  Assessment")
        body_p(doc, c["assessment"])

    # ── 12. Homework ──────────────────────────────────────────────────────────
    if c.get("homework"):
        section_heading(doc, "9.  Homework / Follow-up")
        body_p(doc, c["homework"])

    # ── 13. Differentiation ───────────────────────────────────────────────────
    diff = c.get("differentiation", {})
    if diff:
        section_heading(doc, "10.  Differentiation")
        diff_tbl = doc.add_table(rows=2, cols=2)
        diff_tbl.style = "Table Grid"
        diff_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (label, key, hdr_hex, val_hex) in enumerate([
            ("Support (SEN / EAL)", "support",   "059669", "ECFDF5"),
            ("Extension (G&T)",     "extension", "7C3AED", "F5F3FF"),
        ]):
            hc = diff_tbl.rows[0].cells[i]
            vc = diff_tbl.rows[1].cells[i]
            cell_bg(hc, hdr_hex); cell_bg(vc, val_hex)
            hr_ = hc.paragraphs[0].add_run(label)
            hr_.font.name = "Calibri"; hr_.font.size = Pt(9); hr_.bold = True
            hr_.font.color.rgb = WHITE
            vr_ = vc.paragraphs[0].add_run(diff.get(key, "N/A"))
            vr_.font.name = "Calibri"; vr_.font.size = Pt(10)
        doc.add_paragraph()

    # ── 14. Teacher Reflection (blank lines) ─────────────────────────────────
    section_heading(doc, "11.  Teacher Reflection / Post-Lesson Notes")
    for _ in range(5):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        para_underline(p, color_hex="D1D5DB", sz=4)
    doc.add_paragraph()

    # ── 15. Signature block ───────────────────────────────────────────────────
    sig = doc.add_table(rows=2, cols=3)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (lbl, val) in enumerate(zip(
        ["Teacher's Signature", "Head of Department", "Date"],
        [teacher_name, "", plan.created_at.strftime("%d / %m / %Y")],
    )):
        tc = sig.rows[0].cells[i]; bc2 = sig.rows[1].cells[i]
        tp2 = tc.paragraphs[0]; tp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tr2 = tp2.add_run(val)
        tr2.font.name = "Calibri"; tr2.font.size = Pt(10)
        tr2.font.color.rgb = DARK_BLUE
        cell_top_border(bc2)
        bp = bc2.paragraphs[0]; bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        br_ = bp.add_run(lbl)
        br_.font.name = "Calibri"; br_.font.size = Pt(8)
        br_.font.color.rgb = MID_GRAY

    buf = io.BytesIO()
    doc.save(buf); buf.seek(0)
    safe = plan.title.replace(" ", "_")[:40]
    response = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response["Content-Disposition"] = f'attachment; filename="LessonPlan_{safe}.docx"'
    return response


# ===========================================================================
#  PDF builder — professional, submission-ready
# ===========================================================================

def _build_pdf(plan):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer,
        Table, TableStyle, KeepTogether,
    )

    teacher      = plan.teacher
    school       = getattr(teacher, "school", None)
    school_name  = school.name if school else "School"
    teacher_name = teacher.get_full_name() or teacher.username
    c = plan.content if isinstance(plan.content, dict) else {}

    INDIGO      = colors.HexColor("#4F46E5")
    DARK_BLUE   = colors.HexColor("#1E1B4B")
    LT_INDIGO   = colors.HexColor("#EEF2FF")
    MID_GRAY    = colors.HexColor("#6B7280")
    LIGHT_GRAY  = colors.HexColor("#F9FAFB")
    BORDER_BLUE = colors.HexColor("#E0E7FF")
    GREEN       = colors.HexColor("#059669")
    PURPLE      = colors.HexColor("#7C3AED")

    PAGE_W, PAGE_H = A4
    MARGIN = 2 * cm
    avail_w = PAGE_W - 2 * MARGIN

    buf = io.BytesIO()

    def _page_template(canv, doc):
        canv.saveState()
        banner_h = 2.2 * cm
        canv.setFillColor(DARK_BLUE)
        canv.rect(0, PAGE_H - banner_h, PAGE_W, banner_h, fill=1, stroke=0)
        canv.setFillColor(colors.HexColor("#C7D2FE"))
        canv.setFont("Helvetica", 8)
        canv.drawString(MARGIN, PAGE_H - 0.88 * cm, school_name.upper())
        canv.setFont("Helvetica-Bold", 8)
        canv.drawRightString(PAGE_W - MARGIN, PAGE_H - 0.88 * cm, "LESSON PLAN")
        canv.setFillColor(colors.white)
        canv.setFont("Helvetica-Bold", 11)
        canv.drawCentredString(PAGE_W / 2, PAGE_H - 1.65 * cm, plan.title[:75])
        canv.setFillColor(DARK_BLUE)
        canv.rect(0, 0, PAGE_W, 1.2 * cm, fill=1, stroke=0)
        canv.setFillColor(colors.HexColor("#C7D2FE"))
        canv.setFont("Helvetica", 7.5)
        canv.drawString(MARGIN, 0.43 * cm,
                        f"Teacher: {teacher_name}  |  {school_name}")
        canv.drawRightString(PAGE_W - MARGIN, 0.43 * cm, f"Page {doc.page}")
        canv.restoreState()

    doc_pdf = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=MARGIN, rightMargin=MARGIN,
        topMargin=2.8 * cm, bottomMargin=1.8 * cm,
    )

    def S(name, **kw):
        return ParagraphStyle(name, **kw)

    title_st   = S("T",  fontSize=15, textColor=DARK_BLUE,
                   fontName="Helvetica-Bold", alignment=TA_CENTER,
                   spaceAfter=4, spaceBefore=4)
    sec_hdr_st = S("SH", fontSize=9, textColor=colors.white,
                   fontName="Helvetica-Bold", alignment=TA_LEFT,
                   spaceBefore=0, spaceAfter=0, leading=13)
    body_st    = S("B",  fontSize=9.5, textColor=colors.HexColor("#111827"),
                   fontName="Helvetica", leading=14, spaceBefore=2, spaceAfter=2)
    note_st    = S("N",  fontSize=8.5, textColor=colors.HexColor("#92400E"),
                   fontName="Helvetica-Oblique", leading=11, spaceBefore=2)
    lbl_st     = S("L",  fontSize=8,  textColor=colors.white,
                   fontName="Helvetica-Bold", alignment=TA_CENTER, leading=10)
    val_st     = S("V",  fontSize=9.5, textColor=DARK_BLUE,
                   fontName="Helvetica-Bold", alignment=TA_CENTER, leading=12)
    gray_lbl   = S("GL", fontSize=8,  textColor=MID_GRAY,
                   fontName="Helvetica", alignment=TA_CENTER, leading=10)
    diff_hdr_st= S("DH", fontSize=9,  textColor=colors.white,
                   fontName="Helvetica-Bold", leading=12)

    def section_bar(number, title):
        t = Table([[Paragraph(f"{number}  {title}", sec_hdr_st)]],
                  colWidths=[avail_w])
        t.setStyle(TableStyle([
            ("BACKGROUND",    (0, 0), (-1, -1), INDIGO),
            ("LEFTPADDING",   (0, 0), (-1, -1), 10),
            ("RIGHTPADDING",  (0, 0), (-1, -1), 10),
            ("TOPPADDING",    (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        return t

    def bp(text):
        return Paragraph(f"&bull;&nbsp; {text}", body_st)

    story = []

    # Title
    story.append(Paragraph(plan.title, title_st))
    story.append(Spacer(1, 4))

    # ── Metadata strip ────────────────────────────────────────────────────────
    meta_data = [
        [Paragraph("Subject",       lbl_st), Paragraph("Grade / Class", lbl_st),
         Paragraph("Duration",      lbl_st), Paragraph("Date Prepared", lbl_st)],
        [Paragraph(plan.subject,    val_st), Paragraph(plan.grade,      val_st),
         Paragraph(f"{plan.duration} min",  val_st),
         Paragraph(plan.created_at.strftime("%d %b %Y"), val_st)],
        [Paragraph("Teacher",       lbl_st), Paragraph("School",        lbl_st),
         Paragraph("Curriculum",    lbl_st), Paragraph("",              lbl_st)],
        [Paragraph(teacher_name,    val_st), Paragraph(school_name,     val_st),
         Paragraph(c.get("curriculum", "Standard"), val_st),
         Paragraph("",             val_st)],
    ]
    meta_tbl = Table(meta_data, colWidths=[avail_w / 4] * 4)
    meta_tbl.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, 0), DARK_BLUE),
        ("BACKGROUND",    (0, 2), (-1, 2), DARK_BLUE),
        ("BACKGROUND",    (0, 1), (-1, 1), LT_INDIGO),
        ("BACKGROUND",    (0, 3), (-1, 3), LT_INDIGO),
        ("TOPPADDING",    (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("GRID",          (0, 0), (-1, -1), 0.5, BORDER_BLUE),
        ("ALIGN",         (0, 0), (-1, -1), "CENTER"),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(meta_tbl)
    story.append(Spacer(1, 10))

    # ── Overview ──────────────────────────────────────────────────────────────
    if c.get("overview"):
        story.append(KeepTogether([
            section_bar("1.", "Lesson Overview"),
            Spacer(1, 4),
            Paragraph(c["overview"], body_st),
            Spacer(1, 6),
        ]))

    # ── Prior Knowledge ───────────────────────────────────────────────────────
    if c.get("prior_knowledge"):
        story.append(KeepTogether([
            section_bar("2.", "Prior Knowledge Required"),
            Spacer(1, 4),
            Paragraph(c["prior_knowledge"], body_st),
            Spacer(1, 6),
        ]))

    # ── Learning Objectives ───────────────────────────────────────────────────
    objs = c.get("objectives", [plan.objectives])
    story.append(KeepTogether([
        section_bar("3.", "Learning Objectives"),
        Spacer(1, 4),
        *[bp(o) for o in objs],
        Spacer(1, 6),
    ]))

    # ── Success Criteria ──────────────────────────────────────────────────────
    if c.get("success_criteria"):
        story.append(KeepTogether([
            section_bar("4.", "Success Criteria"),
            Spacer(1, 4),
            *[bp(sc) for sc in c["success_criteria"]],
            Spacer(1, 6),
        ]))

    # ── Key Vocabulary ────────────────────────────────────────────────────────
    if c.get("key_vocabulary"):
        vocab = c["key_vocabulary"]
        vocab_w = avail_w / max(len(vocab), 1)
        vocab_tbl = Table(
            [[Paragraph(v, S("KV", fontSize=9, textColor=DARK_BLUE,
                             fontName="Helvetica-Bold", alignment=TA_CENTER,
                             leading=12)) for v in vocab]],
            colWidths=[vocab_w] * len(vocab),
        )
        vocab_tbl.setStyle(TableStyle([
            ("BACKGROUND",    (0, 0), (-1, -1), LT_INDIGO),
            ("GRID",          (0, 0), (-1, -1), 0.5, BORDER_BLUE),
            ("TOPPADDING",    (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("ALIGN",         (0, 0), (-1, -1), "CENTER"),
        ]))
        story.append(KeepTogether([
            section_bar("5.", "Key Vocabulary"),
            Spacer(1, 4),
            vocab_tbl,
            Spacer(1, 6),
        ]))

    # ── Materials ─────────────────────────────────────────────────────────────
    if c.get("materials"):
        story.append(KeepTogether([
            section_bar("6.", "Materials & Resources"),
            Spacer(1, 4),
            *[bp(m) for m in c["materials"]],
            Spacer(1, 6),
        ]))

    # ── Lesson Procedure ──────────────────────────────────────────────────────
    sections = c.get("sections", [])
    if sections:
        proc_col_w = [2.6 * cm, 1.4 * cm, 4.7 * cm, 4.7 * cm, 2.6 * cm]
        proc_rows = [[
            Paragraph("Phase",              sec_hdr_st),
            Paragraph("Time\n(min)",        sec_hdr_st),
            Paragraph("Teacher Activities", sec_hdr_st),
            Paragraph("Student Activities", sec_hdr_st),
            Paragraph("Notes",              sec_hdr_st),
        ]]
        for s in sections:
            t_acts = s.get("teacher_activities", s.get("activities", []))
            s_acts = s.get("student_activities", [])
            notes  = s.get("teacher_notes", "")
            proc_rows.append([
                Paragraph(f"<b>{s.get('name', '')}</b>", body_st),
                Paragraph(str(s.get("duration", "")), body_st),
                Paragraph("<br/>".join(f"&bull; {a}" for a in t_acts), body_st),
                Paragraph("<br/>".join(f"&bull; {a}" for a in s_acts), body_st),
                (Paragraph(f"<i>{notes}</i>", note_st) if notes
                 else Paragraph("", body_st)),
            ])
        ts = [
            ("BACKGROUND",    (0, 0), (-1, 0), INDIGO),
            ("FONTSIZE",      (0, 0), (-1, -1), 8.5),
            ("ALIGN",         (0, 0), (-1, 0), "CENTER"),
            ("ALIGN",         (1, 1), (1, -1), "CENTER"),
            ("VALIGN",        (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING",    (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LEFTPADDING",   (0, 0), (-1, -1), 5),
            ("RIGHTPADDING",  (0, 0), (-1, -1), 5),
            ("GRID",          (0, 0), (-1, -1), 0.5, BORDER_BLUE),
            ("TEXTCOLOR",     (0, 1), (0, -1), INDIGO),
        ]
        for i in range(1, len(proc_rows)):
            bg = "#FFFFFF" if i % 2 == 1 else "#EEF2FF"
            ts.append(("BACKGROUND", (1, i), (-1, i), colors.HexColor(bg)))
        proc_tbl = Table(proc_rows, colWidths=proc_col_w, repeatRows=1)
        proc_tbl.setStyle(TableStyle(ts))
        story.append(section_bar("7.", "Lesson Procedure"))
        story.append(Spacer(1, 4))
        story.append(proc_tbl)
        story.append(Spacer(1, 6))

    # ── Assessment ────────────────────────────────────────────────────────────
    if c.get("assessment"):
        story.append(KeepTogether([
            section_bar("8.", "Assessment"),
            Spacer(1, 4),
            Paragraph(c["assessment"], body_st),
            Spacer(1, 6),
        ]))

    # ── Homework ──────────────────────────────────────────────────────────────
    if c.get("homework"):
        story.append(KeepTogether([
            section_bar("9.", "Homework / Follow-up"),
            Spacer(1, 4),
            Paragraph(c["homework"], body_st),
            Spacer(1, 6),
        ]))

    # ── Differentiation ───────────────────────────────────────────────────────
    diff = c.get("differentiation", {})
    if diff:
        diff_data = [
            [Paragraph("Support (SEN / EAL)", diff_hdr_st),
             Paragraph("Extension (G&T)",      diff_hdr_st)],
            [Paragraph(diff.get("support",   "N/A"), body_st),
             Paragraph(diff.get("extension", "N/A"), body_st)],
        ]
        diff_tbl = Table(diff_data, colWidths=[avail_w / 2] * 2)
        diff_tbl.setStyle(TableStyle([
            ("BACKGROUND",    (0, 0), (0, 0), GREEN),
            ("BACKGROUND",    (1, 0), (1, 0), PURPLE),
            ("BACKGROUND",    (0, 1), (0, 1), colors.HexColor("#ECFDF5")),
            ("BACKGROUND",    (1, 1), (1, 1), colors.HexColor("#F5F3FF")),
            ("TOPPADDING",    (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("LEFTPADDING",   (0, 0), (-1, -1), 9),
            ("RIGHTPADDING",  (0, 0), (-1, -1), 9),
            ("GRID",          (0, 0), (-1, -1), 0.5, BORDER_BLUE),
            ("VALIGN",        (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(KeepTogether([
            section_bar("10.", "Differentiation"),
            Spacer(1, 4),
            diff_tbl,
            Spacer(1, 6),
        ]))

    # ── Teacher Reflection (blank lines) ─────────────────────────────────────
    refl_rows = [[""] for _ in range(5)]
    refl_tbl  = Table(refl_rows, colWidths=[avail_w], rowHeights=[0.9 * cm] * 5)
    refl_tbl.setStyle(TableStyle([
        ("GRID",       (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
        ("BACKGROUND", (0, 0), (-1, -1), LIGHT_GRAY),
    ]))
    story.append(KeepTogether([
        section_bar("11.", "Teacher Reflection / Post-Lesson Notes"),
        Spacer(1, 4),
        refl_tbl,
        Spacer(1, 10),
    ]))

    # ── Signature row ─────────────────────────────────────────────────────────
    sig_data = [
        [Paragraph("Teacher's Signature", gray_lbl),
         Paragraph("Head of Department",  gray_lbl),
         Paragraph("Date",                gray_lbl)],
        [Paragraph(teacher_name,                              val_st),
         Paragraph("",                                        val_st),
         Paragraph(plan.created_at.strftime("%d / %m / %Y"), val_st)],
    ]
    sig_tbl = Table(sig_data, colWidths=[avail_w / 3] * 3)
    sig_tbl.setStyle(TableStyle([
        ("FONTSIZE",      (0, 0), (-1, -1), 8.5),
        ("ALIGN",         (0, 0), (-1, -1), "CENTER"),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING",    (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LINEABOVE",     (0, 1), (-1, 1), 1.5, DARK_BLUE),
        ("GRID",          (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
    ]))
    story.append(sig_tbl)

    doc_pdf.build(story, onFirstPage=_page_template, onLaterPages=_page_template)
    buf.seek(0)
    safe = plan.title.replace(" ", "_")[:40]
    response = HttpResponse(buf.read(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="LessonPlan_{safe}.pdf"'
    return response


# ---------------------------------------------------------------------------
# Other teacher views
# ---------------------------------------------------------------------------

@role_required("teacher")
def assessment_generator(request):
    return render(request, "teacher/assessment_generator.html")


@role_required("teacher")
def content_generator(request):
    return render(request, "teacher/content_generator.html")


@role_required("teacher")
def teacher_agent(request):
    return render(request, "teacher/teacher_agent.html")


@role_required("teacher")
def attendance(request):
    return render(request, "teacher/attendance.html")


@role_required("teacher")
def timetable(request):
    return render(request, "teacher/timetable.html")


@role_required("teacher")
def students(request):
    return render(request, "teacher/students.html")


def logout_view(request):
    logout(request)
    return redirect("/login/")
